from docx import Document

from langchain_core.documents import Document as LCDocument

from app.loaders.base_loader import BaseLoader


class DOCXLoader(BaseLoader):
    def load(self, file_path: str):
        document = Document(file_path)

        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

        return [
            LCDocument(
                page_content=text,
                metadata={
                    "source": file_path
                }
            )
        ]